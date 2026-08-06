#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Скрипт для сравнения отчётов SAPE
Сравнивает автоматически сгенерированный отчёт с ручным заполнением

Использование:
    python3 compare_reports.py <автоматический_отчёт.xlsx> <ручной_отчёт.xlsx>
"""

import sys
import openpyxl
from openpyxl import load_workbook
from typing import Dict, List, Tuple, Optional
import re


class ReportComparator:
    def __init__(self, auto_report_path: str, manual_report_path: str):
        self.auto_report_path = auto_report_path
        self.manual_report_path = manual_report_path
        self.max_frequency = 3.0  # Максимальная частота для охватов
        self.tolerance = 0.0001  # 0.01% допуск для расхождений

    def load_workbook_sheets(self, path: str) -> Dict[str, openpyxl.worksheet.worksheet.Worksheet]:
        """Загрузить все листы из Excel файла"""
        wb = load_workbook(path, data_only=True)
        return {sheet.title: sheet for sheet in wb.worksheets}

    def find_campaign_sections(self, sheet) -> List[Dict]:
        """Найти все секции кампаний на листе"""
        sections = []

        # Поиск заголовков "Показатели кампании"
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=False), start=1):
            for col_idx, cell in enumerate(row, start=1):
                if cell.value and 'Показатели кампании' in str(cell.value):
                    # Извлечь ID кампании (если есть)
                    campaign_ids = re.findall(r'\b\d{6}\b', str(cell.value))

                    # Если нет ID, использовать название из скобок
                    campaign_name = str(cell.value)
                    if not campaign_ids:
                        # Попробовать извлечь название из скобок, например "(OLV Москва)"
                        match = re.search(r'\(([^)]+)\)', campaign_name)
                        if match:
                            campaign_name = match.group(1).strip()
                        campaign_ids = [campaign_name]  # Использовать название как идентификатор

                    # Найти строку с заголовками колонок (должна быть через 1-2 строки)
                    header_row_idx = None
                    for offset in range(1, 4):
                        check_row = row_idx + offset
                        if check_row <= sheet.max_row:
                            first_cell = sheet.cell(check_row, col_idx).value
                            if first_cell and 'дата' in str(first_cell).lower():
                                header_row_idx = check_row
                                break

                    if header_row_idx:
                        # Определить колонки данных
                        headers = {}
                        for offset in range(10):
                            header_cell = sheet.cell(header_row_idx, col_idx + offset)
                            if header_cell.value:
                                header_name = str(header_cell.value).lower().strip()
                                if 'дата' in header_name:
                                    headers['date'] = col_idx + offset
                                elif 'показы' in header_name or 'shows' in header_name:
                                    headers['impressions'] = col_idx + offset
                                elif 'охват' in header_name or 'reach' in header_name:
                                    headers['reach'] = col_idx + offset
                                elif 'клики' in header_name or 'clicks' in header_name:
                                    headers['clicks'] = col_idx + offset
                                elif 'досмотры' in header_name or 'completes' in header_name:
                                    headers['completes'] = col_idx + offset

                        sections.append({
                            'campaign_ids': campaign_ids,
                            'campaign_name': campaign_name,
                            'header_row': header_row_idx,
                            'data_start_row': header_row_idx + 1,
                            'header_col': col_idx,
                            'columns': headers
                        })

        return sections

    def extract_section_data(self, sheet, section: Dict) -> Dict:
        """Извлечь данные из секции"""
        from datetime import datetime

        data = {
            'daily': [],
            'total': {}
        }

        date_col = section['columns'].get('date')
        impressions_col = section['columns'].get('impressions')
        reach_col = section['columns'].get('reach')
        clicks_col = section['columns'].get('clicks')
        completes_col = section['columns'].get('completes')

        if not date_col:
            return data

        # Читать данные построчно
        for row_idx in range(section['data_start_row'], section['data_start_row'] + 50):
            date_cell = sheet.cell(row_idx, date_col).value

            if not date_cell:
                continue

            # Преобразовать дату в строку
            if isinstance(date_cell, datetime):
                date_str = date_cell.strftime('%d.%m.%Y')
            else:
                date_str = str(date_cell).strip()

            # Проверка на строку "Всего"
            if 'всего' in date_str.lower():
                # Это итоговая строка
                if impressions_col:
                    data['total']['impressions'] = self.get_numeric_value(sheet.cell(row_idx, impressions_col).value)
                if reach_col:
                    data['total']['reach'] = self.get_numeric_value(sheet.cell(row_idx, reach_col).value)
                if clicks_col:
                    data['total']['clicks'] = self.get_numeric_value(sheet.cell(row_idx, clicks_col).value)
                if completes_col:
                    data['total']['completes'] = self.get_numeric_value(sheet.cell(row_idx, completes_col).value)
                break

            # Обычная строка с датой
            daily_data = {'date': date_str}

            if impressions_col:
                daily_data['impressions'] = self.get_numeric_value(sheet.cell(row_idx, impressions_col).value)
            if reach_col:
                daily_data['reach'] = self.get_numeric_value(sheet.cell(row_idx, reach_col).value)
            if clicks_col:
                daily_data['clicks'] = self.get_numeric_value(sheet.cell(row_idx, clicks_col).value)
            if completes_col:
                daily_data['completes'] = self.get_numeric_value(sheet.cell(row_idx, completes_col).value)

            # Пропускаем строки без данных
            if daily_data.get('impressions', 0) > 0 or daily_data.get('clicks', 0) > 0:
                data['daily'].append(daily_data)

        return data

    def get_numeric_value(self, cell_value) -> float:
        """Преобразовать значение ячейки в число"""
        if cell_value is None:
            return 0.0

        if isinstance(cell_value, (int, float)):
            return float(cell_value)

        # Попытка преобразовать строку
        try:
            # Убрать пробелы и запятые
            value_str = str(cell_value).replace(' ', '').replace(',', '')
            return float(value_str)
        except:
            return 0.0

    def compare_values(self, auto_val: float, manual_val: float, metric_name: str) -> Tuple[bool, float, str]:
        """
        Сравнить два значения

        Returns:
            (is_ok, diff_percent, message)
        """
        if manual_val == 0:
            if auto_val == 0:
                return True, 0.0, "OK"
            else:
                return False, 100.0, f"Ручной отчёт: 0, автоматический: {auto_val}"

        diff = abs(auto_val - manual_val)
        diff_percent = (diff / manual_val) * 100

        if metric_name in ['impressions', 'clicks', 'completes']:
            # Для показов/кликов допустимо 0.01% расхождение
            is_ok = diff_percent <= 0.01
            if is_ok:
                return True, diff_percent, f"OK (расхождение {diff_percent:.4f}%)"
            else:
                return False, diff_percent, f"Расхождение {diff_percent:.4f}% (допустимо 0.01%)"

        return True, diff_percent, f"OK ({diff_percent:.4f}%)"

    def validate_reach(self, daily_data: List[Dict], total_reach: float) -> Tuple[bool, str]:
        """
        Проверить корректность охватов

        Требования:
        1. Частота (показы/охват) <= max_frequency для каждого дня
        2. Итоговый охват <= суммы дневных охватов
        """
        issues = []

        # Проверка дневных частот
        for day in daily_data:
            impressions = day.get('impressions', 0)
            reach = day.get('reach', 0)

            if reach > 0 and impressions > 0:
                frequency = impressions / reach
                if frequency > self.max_frequency:
                    issues.append(
                        f"  Дата {day['date']}: частота {frequency:.2f} > {self.max_frequency} "
                        f"(показы: {impressions}, охват: {reach})"
                    )

        # Проверка итогового охвата
        sum_daily_reach = sum(day.get('reach', 0) for day in daily_data)
        if total_reach > sum_daily_reach:
            issues.append(
                f"  Итоговый охват {total_reach} > суммы дневных охватов {sum_daily_reach}"
            )

        if issues:
            return False, "\n".join(issues)
        else:
            return True, "OK"

    def compare_reports(self):
        """Основная функция сравнения отчётов"""
        print("=" * 80)
        print("СРАВНЕНИЕ ОТЧЁТОВ SAPE")
        print("=" * 80)
        print(f"Автоматический: {self.auto_report_path}")
        print(f"Ручной:         {self.manual_report_path}")
        print()

        # Загрузить отчёты
        print("Загрузка отчётов...")
        auto_sheets = self.load_workbook_sheets(self.auto_report_path)
        manual_sheets = self.load_workbook_sheets(self.manual_report_path)

        print(f"Автоматический отчёт: {len(auto_sheets)} листов")
        print(f"Ручной отчёт: {len(manual_sheets)} листов")
        print()

        total_errors = 0
        total_warnings = 0

        # Сравнить каждый лист
        for sheet_name in auto_sheets.keys():
            if sheet_name not in manual_sheets:
                print(f"⚠️  Лист '{sheet_name}' отсутствует в ручном отчёте")
                total_warnings += 1
                continue

            print(f"\n{'=' * 80}")
            print(f"ЛИСТ: {sheet_name}")
            print(f"{'=' * 80}")

            auto_sheet = auto_sheets[sheet_name]
            manual_sheet = manual_sheets[sheet_name]

            # Найти секции кампаний
            auto_sections = self.find_campaign_sections(auto_sheet)
            manual_sections = self.find_campaign_sections(manual_sheet)

            print(f"Найдено секций: автоматический={len(auto_sections)}, ручной={len(manual_sections)}")

            # Сравнить секции по порядку (предполагаем что порядок одинаковый)
            for idx, auto_section in enumerate(auto_sections):
                campaign_id = auto_section['campaign_ids'][0] if auto_section['campaign_ids'] else 'Unknown'

                print(f"\n--- Секция {idx + 1}: {campaign_id} ---")

                # Для файлов без ID сравниваем секции по порядку
                manual_section = None
                if idx < len(manual_sections):
                    manual_section = manual_sections[idx]
                else:
                    print(f"❌ Секция {idx + 1} не найдена в ручном отчёте")
                    total_errors += 1
                    continue

                # Извлечь данные
                auto_data = self.extract_section_data(auto_sheet, auto_section)
                manual_data = self.extract_section_data(manual_sheet, manual_section)

                print(f"Дней с данными: автоматический={len(auto_data['daily'])}, ручной={len(manual_data['daily'])}")

                # Сравнить дневные данные
                section_errors = 0
                for auto_day in auto_data['daily']:
                    date = auto_day['date']

                    # Найти соответствующий день в ручном отчёте
                    manual_day = None
                    for md in manual_data['daily']:
                        if md['date'] == date:
                            manual_day = md
                            break

                    if not manual_day:
                        continue  # День отсутствует в ручном отчёте (возможно ещё не заполнен)

                    # Сравнить показы
                    if 'impressions' in auto_day and 'impressions' in manual_day:
                        is_ok, diff_pct, msg = self.compare_values(
                            auto_day['impressions'],
                            manual_day['impressions'],
                            'impressions'
                        )
                        if not is_ok:
                            print(f"  ❌ {date} - Показы: {msg}")
                            print(f"     Автоматический: {auto_day['impressions']}, Ручной: {manual_day['impressions']}")
                            section_errors += 1

                    # Сравнить клики
                    if 'clicks' in auto_day and 'clicks' in manual_day:
                        is_ok, diff_pct, msg = self.compare_values(
                            auto_day['clicks'],
                            manual_day['clicks'],
                            'clicks'
                        )
                        if not is_ok:
                            print(f"  ❌ {date} - Клики: {msg}")
                            print(f"     Автоматический: {auto_day['clicks']}, Ручной: {manual_day['clicks']}")
                            section_errors += 1

                    # Сравнить досмотры (если есть)
                    if 'completes' in auto_day and 'completes' in manual_day:
                        is_ok, diff_pct, msg = self.compare_values(
                            auto_day['completes'],
                            manual_day['completes'],
                            'completes'
                        )
                        if not is_ok:
                            print(f"  ❌ {date} - Досмотры: {msg}")
                            print(f"     Автоматический: {auto_day['completes']}, Ручной: {manual_day['completes']}")
                            section_errors += 1

                # Проверить охваты
                print(f"\nПроверка охватов:")
                is_reach_ok, reach_msg = self.validate_reach(
                    auto_data['daily'],
                    auto_data['total'].get('reach', 0)
                )
                if is_reach_ok:
                    print(f"  ✅ {reach_msg}")
                else:
                    print(f"  ❌ Проблемы с охватами:")
                    print(reach_msg)
                    section_errors += 1

                # Сравнить итоги
                print(f"\nИтоговые значения:")
                if 'impressions' in auto_data['total'] and 'impressions' in manual_data['total']:
                    is_ok, diff_pct, msg = self.compare_values(
                        auto_data['total']['impressions'],
                        manual_data['total']['impressions'],
                        'impressions'
                    )
                    status = "✅" if is_ok else "❌"
                    print(f"  {status} Показы: {msg}")
                    print(f"     Автоматический: {auto_data['total']['impressions']}, Ручной: {manual_data['total']['impressions']}")
                    if not is_ok:
                        section_errors += 1

                if 'clicks' in auto_data['total'] and 'clicks' in manual_data['total']:
                    is_ok, diff_pct, msg = self.compare_values(
                        auto_data['total']['clicks'],
                        manual_data['total']['clicks'],
                        'clicks'
                    )
                    status = "✅" if is_ok else "❌"
                    print(f"  {status} Клики: {msg}")
                    print(f"     Автоматический: {auto_data['total']['clicks']}, Ручной: {manual_data['total']['clicks']}")
                    if not is_ok:
                        section_errors += 1

                if section_errors == 0:
                    print(f"\n✅ Кампания {campaign_id}: Все проверки пройдены")
                else:
                    print(f"\n❌ Кампания {campaign_id}: Найдено {section_errors} ошибок")
                    total_errors += section_errors

        # Итоговый результат
        print(f"\n\n{'=' * 80}")
        print("ИТОГОВЫЙ РЕЗУЛЬТАТ")
        print(f"{'=' * 80}")

        if total_errors == 0 and total_warnings == 0:
            print("✅ Все проверки пройдены успешно!")
        else:
            print(f"❌ Найдено ошибок: {total_errors}")
            print(f"⚠️  Найдено предупреждений: {total_warnings}")

        return total_errors == 0


def main():
    if len(sys.argv) != 3:
        print("Использование: python3 compare_reports.py <автоматический_отчёт.xlsx> <ручной_отчёт.xlsx>")
        sys.exit(1)

    auto_report = sys.argv[1]
    manual_report = sys.argv[2]

    comparator = ReportComparator(auto_report, manual_report)
    success = comparator.compare_reports()

    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
