#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Детальное сравнение листа "апрель" из двух отчётов
"""

import sys
import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import re


def get_numeric_value(cell_value):
    """Преобразовать значение ячейки в число"""
    if cell_value is None:
        return 0.0
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    try:
        value_str = str(cell_value).replace(' ', '').replace(',', '')
        return float(value_str)
    except:
        return 0.0


def find_campaign_sections(sheet):
    """Найти все секции кампаний на листе"""
    sections = []

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=False), start=1):
        for col_idx, cell in enumerate(row, start=1):
            if cell.value and 'Показатели кампании' in str(cell.value):
                campaign_ids = re.findall(r'\b\d{6}\b', str(cell.value))
                campaign_name = str(cell.value)

                if not campaign_ids:
                    match = re.search(r'\(([^)]+)\)', campaign_name)
                    if match:
                        campaign_name = match.group(1).strip()
                    campaign_ids = [campaign_name]

                # Найти строку с заголовками
                header_row_idx = None
                for offset in range(1, 4):
                    check_row = row_idx + offset
                    if check_row <= sheet.max_row:
                        first_cell = sheet.cell(check_row, col_idx).value
                        if first_cell and 'дата' in str(first_cell).lower():
                            header_row_idx = check_row
                            break

                if header_row_idx:
                    headers = {}
                    for offset in range(10):
                        header_cell = sheet.cell(header_row_idx, col_idx + offset)
                        if header_cell.value:
                            header_name = str(header_cell.value).lower().strip()
                            if 'дата' in header_name:
                                headers['date'] = col_idx + offset
                            elif 'показы' in header_name:
                                headers['impressions'] = col_idx + offset
                            elif 'охват' in header_name:
                                headers['reach'] = col_idx + offset
                            elif 'клики' in header_name:
                                headers['clicks'] = col_idx + offset
                            elif 'досмотры' in header_name:
                                headers['completes'] = col_idx + offset

                    sections.append({
                        'campaign_ids': campaign_ids,
                        'campaign_name': campaign_name,
                        'header_row': header_row_idx,
                        'data_start_row': header_row_idx + 1,
                        'columns': headers
                    })

    return sections


def extract_section_data(sheet, section):
    """Извлечь данные из секции"""
    data = {'daily': [], 'total': {}}

    date_col = section['columns'].get('date')
    impressions_col = section['columns'].get('impressions')
    reach_col = section['columns'].get('reach')
    clicks_col = section['columns'].get('clicks')
    completes_col = section['columns'].get('completes')

    if not date_col:
        return data

    for row_idx in range(section['data_start_row'], section['data_start_row'] + 50):
        date_cell = sheet.cell(row_idx, date_col).value

        if not date_cell:
            continue

        if isinstance(date_cell, datetime):
            date_str = date_cell.strftime('%d.%m.%Y')
        else:
            date_str = str(date_cell).strip()

        if 'всего' in date_str.lower():
            if impressions_col:
                data['total']['impressions'] = get_numeric_value(sheet.cell(row_idx, impressions_col).value)
            if reach_col:
                data['total']['reach'] = get_numeric_value(sheet.cell(row_idx, reach_col).value)
            if clicks_col:
                data['total']['clicks'] = get_numeric_value(sheet.cell(row_idx, clicks_col).value)
            if completes_col:
                data['total']['completes'] = get_numeric_value(sheet.cell(row_idx, completes_col).value)
            break

        daily_data = {'date': date_str}

        if impressions_col:
            daily_data['impressions'] = get_numeric_value(sheet.cell(row_idx, impressions_col).value)
        if reach_col:
            daily_data['reach'] = get_numeric_value(sheet.cell(row_idx, reach_col).value)
        if clicks_col:
            daily_data['clicks'] = get_numeric_value(sheet.cell(row_idx, clicks_col).value)
        if completes_col:
            daily_data['completes'] = get_numeric_value(sheet.cell(row_idx, completes_col).value)

        if daily_data.get('impressions', 0) > 0 or daily_data.get('clicks', 0) > 0:
            data['daily'].append(daily_data)

    return data


def compare_april_reports(auto_file, manual_file):
    """Сравнить лист апрель из двух файлов"""

    # Загрузить файлы
    auto_wb = load_workbook(auto_file, data_only=True)
    manual_wb = load_workbook(manual_file, data_only=True)

    auto_sheet = auto_wb['апрель']
    manual_sheet = manual_wb['апрель']

    # Найти секции
    auto_sections = find_campaign_sections(auto_sheet)
    manual_sections = find_campaign_sections(manual_sheet)

    print("=" * 120)
    print(f"СРАВНЕНИЕ ЛИСТА 'АПРЕЛЬ'")
    print("=" * 120)
    print(f"Копия Акира Оил (автомат): {len(auto_sections)} секций")
    print(f"Акира Оил (ручной):        {len(manual_sections)} секций")
    print()

    # Сравним секции по порядку (предполагаем что порядок одинаковый)
    for idx, auto_section in enumerate(auto_sections):
        # Извлечь ID из автоматического отчёта
        campaign_ids = auto_section['campaign_ids']
        campaign_id = campaign_ids[0] if campaign_ids else f'Section{idx+1}'

        print(f"\n{'=' * 120}")
        print(f"СЕКЦИЯ {idx+1}: Кампания {campaign_id}")
        print(f"  Автомат: {auto_section['campaign_name']}")

        # Взять соответствующую секцию из ручного отчёта по порядку
        if idx >= len(manual_sections):
            print(f"⚠️  Секция {idx+1} отсутствует в ручном отчёте")
            continue

        manual_section = manual_sections[idx]
        print(f"  Ручной:  {manual_section['campaign_name']}")
        print(f"{'=' * 120}")

        # Извлечь данные
        auto_data = extract_section_data(auto_sheet, auto_section)
        manual_data = extract_section_data(manual_sheet, manual_section)

        print(f"\nДней с данными: Автомат={len(auto_data['daily'])}, Ручной={len(manual_data['daily'])}")

        # Создать маппинг дат для ручного отчёта
        manual_by_date = {d['date']: d for d in manual_data['daily']}

        # Заголовок таблицы
        print(f"\n{'Дата':<12} | {'Метрика':<10} | {'Автомат (Копия)':>15} | {'Ручной (Оригинал)':>15} | {'Разница':>10} | {'%':>8}")
        print("-" * 120)

        has_issues = False

        # Сравнить дневные данные
        for auto_day in auto_data['daily']:
            date = auto_day['date']
            manual_day = manual_by_date.get(date)

            if not manual_day:
                print(f"{date:<12} | {'ВСЕ':<10} | {'ЕСТЬ':>15} | {'НЕТ ДАННЫХ':>15} | {'-':>10} | {'-':>8}")
                has_issues = True
                continue

            # Сравнить показы
            auto_imp = auto_day.get('impressions', 0)
            manual_imp = manual_day.get('impressions', 0)
            diff_imp = auto_imp - manual_imp
            diff_pct_imp = (abs(diff_imp) / manual_imp * 100) if manual_imp > 0 else 0

            status_imp = "✅" if diff_pct_imp <= 0.01 else "❌"

            print(f"{date:<12} | {'Показы':<10} | {auto_imp:>15,.0f} | {manual_imp:>15,.0f} | {diff_imp:>+10,.0f} | {diff_pct_imp:>7.2f}% {status_imp}")

            if diff_pct_imp > 0.01:
                has_issues = True

            # Сравнить клики
            auto_clicks = auto_day.get('clicks', 0)
            manual_clicks = manual_day.get('clicks', 0)
            diff_clicks = auto_clicks - manual_clicks
            diff_pct_clicks = (abs(diff_clicks) / manual_clicks * 100) if manual_clicks > 0 else 0

            status_clicks = "✅" if diff_pct_clicks <= 0.01 else "❌"

            print(f"{'':<12} | {'Клики':<10} | {auto_clicks:>15,.0f} | {manual_clicks:>15,.0f} | {diff_clicks:>+10,.0f} | {diff_pct_clicks:>7.2f}% {status_clicks}")

            if diff_pct_clicks > 0.01:
                has_issues = True

            # Сравнить досмотры (если есть)
            if 'completes' in auto_day and 'completes' in manual_day:
                auto_comp = auto_day.get('completes', 0)
                manual_comp = manual_day.get('completes', 0)
                diff_comp = auto_comp - manual_comp
                diff_pct_comp = (abs(diff_comp) / manual_comp * 100) if manual_comp > 0 else 0

                status_comp = "✅" if diff_pct_comp <= 0.01 else "❌"

                print(f"{'':<12} | {'Досмотры':<10} | {auto_comp:>15,.0f} | {manual_comp:>15,.0f} | {diff_comp:>+10,.0f} | {diff_pct_comp:>7.2f}% {status_comp}")

                if diff_pct_comp > 0.01:
                    has_issues = True

            # Проверить охват (частота)
            auto_reach = auto_day.get('reach', 0)
            if auto_reach > 0 and auto_imp > 0:
                frequency = auto_imp / auto_reach
                if frequency > 3.0:
                    print(f"{'':<12} | {'⚠️  Охват':<10} | {auto_reach:>15,.0f} | {'частота':>15} | {frequency:>10.2f} | {'> 3.0':>8} ❌")
                    has_issues = True

            print("-" * 120)

        # Итоговые значения
        print(f"\n{'ИТОГО':<12} | {'Метрика':<10} | {'Автомат':>15} | {'Ручной':>15} | {'Разница':>10} | {'%':>8}")
        print("-" * 120)

        # Показы итого
        auto_total_imp = auto_data['total'].get('impressions', 0)
        manual_total_imp = manual_data['total'].get('impressions', 0)
        diff_total_imp = auto_total_imp - manual_total_imp
        diff_pct_total_imp = (abs(diff_total_imp) / manual_total_imp * 100) if manual_total_imp > 0 else 0
        status_total_imp = "✅" if diff_pct_total_imp <= 0.01 else "❌"

        print(f"{'ВСЕГО':<12} | {'Показы':<10} | {auto_total_imp:>15,.0f} | {manual_total_imp:>15,.0f} | {diff_total_imp:>+10,.0f} | {diff_pct_total_imp:>7.2f}% {status_total_imp}")

        # Клики итого
        auto_total_clicks = auto_data['total'].get('clicks', 0)
        manual_total_clicks = manual_data['total'].get('clicks', 0)
        diff_total_clicks = auto_total_clicks - manual_total_clicks
        diff_pct_total_clicks = (abs(diff_total_clicks) / manual_total_clicks * 100) if manual_total_clicks > 0 else 0
        status_total_clicks = "✅" if diff_pct_total_clicks <= 0.01 else "❌"

        print(f"{'ВСЕГО':<12} | {'Клики':<10} | {auto_total_clicks:>15,.0f} | {manual_total_clicks:>15,.0f} | {diff_total_clicks:>+10,.0f} | {diff_pct_total_clicks:>7.2f}% {status_total_clicks}")

        # Досмотры итого
        if 'completes' in auto_data['total'] and 'completes' in manual_data['total']:
            auto_total_comp = auto_data['total'].get('completes', 0)
            manual_total_comp = manual_data['total'].get('completes', 0)
            diff_total_comp = auto_total_comp - manual_total_comp
            diff_pct_total_comp = (abs(diff_total_comp) / manual_total_comp * 100) if manual_total_comp > 0 else 0
            status_total_comp = "✅" if diff_pct_total_comp <= 0.01 else "❌"

            print(f"{'ВСЕГО':<12} | {'Досмотры':<10} | {auto_total_comp:>15,.0f} | {manual_total_comp:>15,.0f} | {diff_total_comp:>+10,.0f} | {diff_pct_total_comp:>7.2f}% {status_total_comp}")

        # Охват итого
        auto_total_reach = auto_data['total'].get('reach', 0)
        manual_total_reach = manual_data['total'].get('reach', 0)

        # Проверка: итоговый охват <= суммы дневных
        sum_daily_reach = sum(d.get('reach', 0) for d in auto_data['daily'])
        reach_check = "✅" if auto_total_reach <= sum_daily_reach else "❌"

        print(f"{'ВСЕГО':<12} | {'Охват':<10} | {auto_total_reach:>15,.0f} | {'сумма дней':>15} | {sum_daily_reach:>10,.0f} | {reach_check:>8}")

        # Частота итого
        if auto_total_reach > 0 and auto_total_imp > 0:
            total_frequency = auto_total_imp / auto_total_reach
            freq_check = "✅" if total_frequency <= 3.0 else "❌"
            print(f"{'ВСЕГО':<12} | {'Частота':<10} | {total_frequency:>15.2f} | {'<= 3.0':>15} | {'-':>10} | {freq_check:>8}")

        if not has_issues:
            print(f"\n✅ Кампания {campaign_id}: Все проверки пройдены")
        else:
            print(f"\n❌ Кампания {campaign_id}: Обнаружены расхождения")


if __name__ == '__main__':
    auto_file = "/Users/annabereznyak/Downloads/Копия Акира Оил.xlsx"
    manual_file = "/Users/annabereznyak/Downloads/Акира Оил.xlsx"

    compare_april_reports(auto_file, manual_file)
